# -*- coding: utf-8 -*-
"""Incubation チーム OKR（Q3振り返り & Q4）リーダー定例用 ミニマルスライド生成"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- パレット（ミニマル・木のぬくもりを感じる温かみのあるアクセント）----
INK    = RGBColor(0x22, 0x22, 0x22)   # 主要テキスト（ニアブラック）
MUTED  = RGBColor(0x70, 0x70, 0x70)   # 補足テキスト
FAINT  = RGBColor(0xA8, 0xA8, 0xA8)   # ごく薄い
LINE   = RGBColor(0xE4, 0xE1, 0xDC)   # 罫線
ACCENT = RGBColor(0xB8, 0x77, 0x2E)   # 温かみのあるオーカー
ACC_BG = RGBColor(0xF7, 0xF2, 0xEA)   # アクセントの淡い背景
GREEN  = RGBColor(0x4E, 0x8A, 0x5C)   # 達成
AMBER  = RGBColor(0xC4, 0x96, 0x2E)   # 進行中
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PAPER  = RGBColor(0xFD, 0xFC, 0xFA)   # ごくわずかに温かい白

FONT = "Yu Gothic"

EMU = 914400
SW, SH = 13.333, 7.5  # 16:9

prs = Presentation()
prs.slide_width  = Emu(int(SW * EMU))
prs.slide_height = Emu(int(SH * EMU))
BLANK = prs.slide_layouts[6]

MX = 0.92  # 左右マージン


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def setp(p, text, size, color=INK, bold=False, font=FONT, space_after=6,
         line=1.12, align=PP_ALIGN.LEFT):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    try:
        p.line_spacing = line
    except Exception:
        pass
    for r in p.runs:
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        # CJK フォント指定
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', font)
    return p


def rect(s, x, y, w, h, fill=None, line_color=None, line_w=0.75):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color; sp.line.width = Pt(line_w)
    return sp


def accent_tab(s, y=0.62, x=MX, w=0.42):
    rect(s, x, y, w, 0.085, fill=ACCENT)


def header(s, kicker, title, title_size=33):
    accent_tab(s)
    _, tf = box(s, MX, 0.80, SW - 2 * MX, 0.42)
    setp(tf.paragraphs[0], kicker, 13.5, ACCENT, bold=True, space_after=0)
    _, tf2 = box(s, MX, 1.18, SW - 2 * MX, 0.95)
    setp(tf2.paragraphs[0], title, title_size, INK, bold=True, line=1.05)


def footer(s, page):
    _, tf = box(s, MX, SH - 0.5, SW - 2 * MX, 0.3)
    p = tf.paragraphs[0]
    setp(p, "Incubation Team OKR ・ リーダー定例 2026-06-28", 9, FAINT, space_after=0)
    _, tf2 = box(s, SW - MX - 1.0, SH - 0.5, 1.0, 0.3)
    setp(tf2.paragraphs[0], str(page), 9, FAINT, align=PP_ALIGN.RIGHT, space_after=0)


# =========================================================
# 1. タイトル
# =========================================================
s = slide()
rect(s, 0, 0, 0.22, SH, fill=ACCENT)
_, tf = box(s, MX + 0.15, 2.55, SW - 2 * MX, 0.5)
setp(tf.paragraphs[0], "INCUBATION TEAM", 15, ACCENT, bold=True, space_after=4)
_, tf = box(s, MX + 0.15, 3.05, SW - 2 * MX, 1.7)
setp(tf.paragraphs[0], "OKR レビュー", 52, INK, bold=True, space_after=2, line=1.0)
setp(tf.add_paragraph(), "Q3 振り返り ＆ Q4（7月〜）に向けて", 24, MUTED, space_after=0, line=1.1)
rect(s, MX + 0.17, 5.35, 2.0, 0.05, fill=LINE)
_, tf = box(s, MX + 0.15, 5.55, SW - 2 * MX, 0.5)
setp(tf.paragraphs[0], "リーダー定例用サマリー　／　2026-06-28", 14, MUTED, space_after=0)

# =========================================================
# 2. Q3 振り返り サマリー
# =========================================================
s = slide()
header(s, "Q3 ｜ 振り返り", "Objective：来期以降、民主化を推進する準備をする", title_size=27)

cards = [
    ("KR1", "パイプライン増\nCランク以上で 2.35億", "Cランク 3.4億", "145% 達成", GREEN, ACC_BG),
    ("KR2", "AIによる\n業務改善施策", "OS 50% / BIZ 20%", "進行中", AMBER, WHITE),
    ("KR3", "事業戦略策定 ＆\n民主化インパクト指標", "事業戦略策定 100%", "達成", GREEN, ACC_BG),
]
cy = 2.55
cw = (SW - 2 * MX - 2 * 0.4) / 3
for i, (kr, goal, result, status, scolor, bgc) in enumerate(cards):
    cx = MX + i * (cw + 0.4)
    rect(s, cx, cy, cw, 3.35, fill=bgc, line_color=LINE, line_w=1)
    rect(s, cx, cy, cw, 0.09, fill=scolor)
    _, tf = box(s, cx + 0.28, cy + 0.32, cw - 0.56, 2.9)
    setp(tf.paragraphs[0], kr, 22, INK, bold=True, space_after=8)
    setp(tf.add_paragraph(), goal, 13.5, MUTED, line=1.2, space_after=14)
    setp(tf.add_paragraph(), "実績", 10.5, FAINT, bold=True, space_after=2)
    setp(tf.add_paragraph(), result, 15, INK, bold=True, line=1.15, space_after=12)
    # ステータスバッジ
    badge = rect(s, cx + 0.28, cy + 2.78, cw - 1.0, 0.42, fill=scolor)
    _, btf = box(s, cx + 0.28, cy + 2.82, cw - 1.0, 0.36)
    btf.paragraphs[0].alignment = PP_ALIGN.CENTER
    setp(btf.paragraphs[0], status, 12.5, WHITE, bold=True, align=PP_ALIGN.CENTER, space_after=0)
footer(s, 2)

# =========================================================
# 3. KR1
# =========================================================
s = slide()
header(s, "Q3 ｜ KR1", "パイプライン増　145% 達成", title_size=30)
# 大きな実績表示
rect(s, MX, 2.45, 4.4, 3.3, fill=ACC_BG, line_color=LINE, line_w=1)
_, tf = box(s, MX + 0.4, 2.95, 3.6, 2.4)
setp(tf.paragraphs[0], "目標 Cランク以上 2.35億", 13, MUTED, space_after=14)
setp(tf.add_paragraph(), "3.4億", 56, ACCENT, bold=True, space_after=0, line=1.0)
setp(tf.add_paragraph(), "Cランク実績　＝　145% 達成", 14, INK, bold=True, space_after=0)

_, tf = box(s, MX + 5.0, 2.65, SW - 2 * MX - 5.0, 3.0)
setp(tf.paragraphs[0], "レビュー", 14, ACCENT, bold=True, space_after=10)
for t in ["目標 2.35億 を大きく上回り達成。", "補助金案件が新たに入ってきたことが寄与。"]:
    p = tf.add_paragraph()
    setp(p, "・ " + t, 16, INK, line=1.3, space_after=10)
footer(s, 3)

# =========================================================
# 4. KR2
# =========================================================
s = slide()
header(s, "Q3 ｜ KR2", "AIによる業務改善施策　＜進行中＞", title_size=29)

# OS カラム
def kr2_col(x, w, label, pct, pct_color, items):
    rect(s, x, 2.5, w, 3.4, fill=WHITE, line_color=LINE, line_w=1)
    rect(s, x, 2.5, 0.08, 3.4, fill=pct_color)
    _, tf = box(s, x + 0.34, 2.78, w - 0.6, 3.0)
    p0 = tf.paragraphs[0]
    setp(p0, label + "　", 17, INK, bold=True, space_after=0)
    r = p0.add_run(); r.text = pct
    r.font.name = FONT; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = pct_color
    setp(tf.add_paragraph(), "", 4, MUTED, space_after=2)
    for it, is_issue in items:
        p = tf.add_paragraph()
        mark = "△ " if is_issue else "・ "
        setp(p, mark + it, 13, (ACCENT if is_issue else INK), line=1.28, space_after=7,
             bold=is_issue)

colw = (SW - 2 * MX - 0.5) / 2
kr2_col(MX, colw, "OS", "50%", AMBER, [
    ("サポートのボットを作成、運用中", False),
    ("課題：情報の精度が悪い → 改善したい", True),
])
kr2_col(MX + colw + 0.5, colw, "BIZ", "20%", AMBER, [
    ("Claudeで事例収集→リテンションを試行中", False),
    ("VUILD事例は豊富だが、リアルなオーナー事例が必要", False),
    ("課題：事例集め（行脚・SNS）と保存場所・フォーマット・ルール", True),
    ("オーナー事例はClaudeで集めにくい（最初の雛形集めは戸倉）", False),
])
footer(s, 4)

# =========================================================
# 5. KR3
# =========================================================
s = slide()
header(s, "Q3 ｜ KR3", "事業戦略策定 ＆ 民主化インパクト指標　＜達成＞", title_size=26)
rect(s, MX, 2.7, SW - 2 * MX, 2.7, fill=ACC_BG, line_color=LINE, line_w=1)
_, tf = box(s, MX + 0.5, 3.05, SW - 2 * MX - 1.0, 2.2)
p0 = tf.paragraphs[0]
setp(p0, "事業戦略策定済み　", 20, INK, bold=True, space_after=0)
r = p0.add_run(); r.text = "100%"
r.font.name = FONT; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = ACCENT
setp(tf.add_paragraph(), "全社でインパクト指標が出てきた。", 16, MUTED, space_after=0, line=1.3)
footer(s, 5)

# =========================================================
# 6. Q4 全社OKR
# =========================================================
s = slide()
header(s, "Q4（7月〜）｜ 全社の方向性", "基盤を整える（成長を目指しつつ、守備力を強化）", title_size=27)
items = [
    ("KR1", "人事制度の整備"),
    ("KR2", "採用の強化"),
    ("KR3", "業務の効率化"),
]
cw = (SW - 2 * MX - 2 * 0.4) / 3
for i, (kr, t) in enumerate(items):
    cx = MX + i * (cw + 0.4)
    rect(s, cx, 2.55, cw, 1.85, fill=WHITE, line_color=LINE, line_w=1)
    _, tf = box(s, cx + 0.3, 2.8, cw - 0.6, 1.4)
    setp(tf.paragraphs[0], kr, 15, ACCENT, bold=True, space_after=8)
    setp(tf.add_paragraph(), t, 18, INK, bold=True, line=1.2, space_after=0)
# 営業
rect(s, MX, 4.75, SW - 2 * MX, 1.15, fill=ACC_BG, line_color=LINE, line_w=1)
_, tf = box(s, MX + 0.3, 5.0, SW - 2 * MX - 0.6, 0.7)
p0 = tf.paragraphs[0]
setp(p0, "営業　", 15, ACCENT, bold=True, space_after=0)
r = p0.add_run(); r.text = "ユーザーのリテンション施策を打つ"
r.font.name = FONT; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = INK
footer(s, 6)

# =========================================================
# 7. Platform OKR（Q4 案）
# =========================================================
s = slide()
header(s, "Q4（7月〜）｜ Platform OKR 案",
       "Objective：来期に民主化を推進するための下準備・仕掛けづくり", title_size=22)
# KR1
rect(s, MX, 2.45, SW - 2 * MX, 1.15, fill=WHITE, line_color=LINE, line_w=1)
rect(s, MX, 2.45, 0.08, 1.15, fill=ACCENT)
_, tf = box(s, MX + 0.34, 2.62, SW - 2 * MX - 0.6, 0.9)
setp(tf.paragraphs[0], "KR1　来期ヨミ C → B へのアップ", 18, INK, bold=True, space_after=4)
setp(tf.add_paragraph(), "指標：Shopbot 導入台数", 13, MUTED, space_after=0)
# KR2
rect(s, MX, 3.78, SW - 2 * MX, 2.55, fill=ACC_BG, line_color=LINE, line_w=1)
rect(s, MX, 3.78, 0.08, 2.55, fill=ACCENT)
_, tf = box(s, MX + 0.34, 3.95, SW - 2 * MX - 0.6, 2.3)
setp(tf.paragraphs[0], "KR2　ポータルサイト「ShopBot＋」をリリース（導入300台に合わせて）",
     18, INK, bold=True, space_after=5)
setp(tf.add_paragraph(), "指標：地域クリエイター拠点数 ／ ものづくりに触れた人", 13, MUTED, space_after=9)
for t in ["地域クリエイター拠点を ＿＿ 拠点（目標数・定義を決めていく）",
          "★ 地域クリエイターを育てる・フィールドになる拠点",
          "A（挑戦）：大学生向けテストプログラム（木匠塾企画）"]:
    setp(tf.add_paragraph(), "・ " + t, 14, INK, line=1.25, space_after=6)
footer(s, 7)

# =========================================================
# 8. ShopBot+ 提供物
# =========================================================
s = slide()
header(s, "Q4 ｜ ShopBot＋", "「ShopBot＋」で提供するもの", title_size=30)
_, tf = box(s, MX, 1.95, SW - 2 * MX, 0.4)
setp(tf.paragraphs[0], "ツールは、これらを実現するための「手段」", 13, MUTED, space_after=0)

rows = [
    ("教育プログラム", "みき・ふく", ""),
    ("人（人材育成）", "戸倉・もえさん・あおやなぎ", ""),
    ("EMARF", "戸倉・COCO", "★ 要検討"),
    ("製材機・乾燥機", "井上・すずこう", "準備中"),
    ("テンプレ（NESTING的）", "戸倉", "案"),
]
ty = 2.55
rh = 0.72
colx = [MX + 0.3, MX + 4.4, MX + 8.6]
# ヘッダ行
rect(s, MX, ty, SW - 2 * MX, 0.5, fill=INK)
heads = ["提供物", "担当", "状態"]
for i, h in enumerate(heads):
    _, hb = box(s, colx[i], ty + 0.11, 3.8, 0.34)
    setp(hb.paragraphs[0], h, 12.5, WHITE, bold=True, space_after=0)
ty += 0.5
for j, (a, b, c) in enumerate(rows):
    bgc = WHITE if j % 2 == 0 else ACC_BG
    rect(s, MX, ty + j * rh, SW - 2 * MX, rh, fill=bgc, line_color=LINE, line_w=0.5)
    _, b1 = box(s, colx[0], ty + j * rh + 0.18, 3.9, 0.5)
    setp(b1.paragraphs[0], a, 15, INK, bold=True, space_after=0)
    _, b2 = box(s, colx[1], ty + j * rh + 0.2, 4.0, 0.5)
    setp(b2.paragraphs[0], b, 13.5, MUTED, space_after=0)
    _, b3 = box(s, colx[2], ty + j * rh + 0.2, 3.4, 0.5)
    setp(b3.paragraphs[0], c, 13.5, ACCENT if c else MUTED, bold=bool(c), space_after=0)
footer(s, 8)

# =========================================================
# 9. 地域クリエイター拠点の定義
# =========================================================
s = slide()
header(s, "Q4 ｜ 論点", "地域クリエイター拠点の「定義」（検討中）", title_size=28)
# 定義の核
rect(s, MX, 2.45, 5.6, 3.5, fill=ACC_BG, line_color=LINE, line_w=1)
_, tf = box(s, MX + 0.35, 2.72, 5.0, 3.1)
setp(tf.paragraphs[0], "定義の核", 14, ACCENT, bold=True, space_after=10)
for t in ["シェア工房である", "エコシステムがある（メンバーコミュニティ）",
          "TO C・地域住民へWS提供／ものづくりに触れる機会"]:
    setp(tf.add_paragraph(), "★ " + t, 15, INK, line=1.3, space_after=10)

_, tf = box(s, MX + 6.0, 2.5, SW - 2 * MX - 6.0, 3.4)
setp(tf.paragraphs[0], "スタンス・論点", 14, ACCENT, bold=True, space_after=10)
for t in ["工房をオープンにしている（WS実施）とわかりやすい",
          "地域を面白くする・インキュベートする人がいる状態",
          "半クローズの拠点もある → ジャンル分けが必要？",
          "拠点側のメリットは？（要検討）"]:
    setp(tf.add_paragraph(), "・ " + t, 14, INK, line=1.28, space_after=9)
_, tf = box(s, MX, 6.05, SW - 2 * MX, 0.5)
setp(tf.paragraphs[0],
     "候補：松川Pukto／FabLab南小国／花巻／久米製材／Andforest／森町／上松／新庄村／DIYSTUDIO／太宰府／カインズ／小菅村／iti-setouchi　ほか",
     10.5, FAINT, line=1.2, space_after=0)
footer(s, 9)

# =========================================================
# 10. ネクストアクション
# =========================================================
s = slide()
header(s, "まとめ ｜ ネクストアクション", "次に決める・動かすこと", title_size=30)
actions = [
    ("KR2 / OS", "サポートボットの情報精度を改善する"),
    ("KR2 / BIZ", "オーナー事例の集め方と、保存場所・フォーマット・ルールを確定"),
    ("Platform / KR2", "地域クリエイター拠点の「定義」と「目標拠点数」を確定"),
    ("Platform / B", "SB4・自動化のロードマップを別途相談（戸倉・中澤）"),
    ("ShopBot＋", "EMARF まわりの方針を検討"),
]
ay = 2.4
ah = 0.82
for i, (tag, t) in enumerate(actions):
    y = ay + i * ah
    rect(s, MX, y, SW - 2 * MX, ah - 0.12, fill=WHITE, line_color=LINE, line_w=1)
    # チェックボックス
    rect(s, MX + 0.28, y + 0.22, 0.26, 0.26, fill=None, line_color=ACCENT, line_w=1.5)
    _, tagb = box(s, MX + 0.78, y + 0.13, 2.5, 0.5)
    setp(tagb.paragraphs[0], tag, 12.5, ACCENT, bold=True, space_after=0)
    _, tb = box(s, MX + 3.4, y + 0.11, SW - 2 * MX - 3.7, 0.55)
    setp(tb.paragraphs[0], t, 15, INK, line=1.15, space_after=0)
footer(s, 10)

out = "/home/user/claude_test/incubation-okr-q3-q4.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
